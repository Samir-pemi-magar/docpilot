from pathlib import Path

from openpyxl import load_workbook

from app.services.extractors.base import BaseExtractor


class XlsxExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> str:
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            lines = []

            for worksheet in workbook.worksheets:
                lines.append(f"Sheet: {worksheet.title}")

                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        lines.append(" | ".join(values))

            return "\n".join(lines)
        finally:
            workbook.close()
